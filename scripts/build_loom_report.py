"""Build the two-page Folio Loom recording briefing."""

# ruff: noqa: E501, E702

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Folio Loom Recording Brief.docx"
INK = "173239"
TEAL = "285C62"
AMBER = "A56A16"
PALE = "EDF3F2"
GRAY = "D9E1E1"


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    properties.append(element)


def borders(table, color=GRAY):
    properties = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        element.append(node)
    properties.append(element)


def cell_text(cell, text, bold=False, color=None, size=9.4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def para(doc, text, size=10.2, color=INK, bold=False, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.06
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    for run in p.runs:
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(INK)
    p.add_run(text).font.size = Pt(9.5)


def setup_section(section):
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("FOLIO  /  LOOM RECORDING BRIEF")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Local sales demo  •  Synthetic documents only  •  September 2026")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("66777A")


def build():
    doc = Document()
    setup_section(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Folio Loom Recording Brief")
    r.font.name = "Aptos Display"
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string("000000")
    subtitle = para(
        doc,
        "A two-minute setup and a 60-second proof of reliable document automation",
        11,
        TEAL,
        space_after=10,
    )
    subtitle.runs[0].italic = True

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(2.05)
    table.columns[1].width = Inches(2.05)
    table.columns[2].width = Inches(2.05)
    borders(table)
    facts = [
        ("Primary path", "Groq Qwen 3.8 27B", "Groq live extraction"),
        ("Guardrail", "Decimal verification", "Model cannot approve itself"),
        ("Recovery", "Local Qwen 3 VL 4B", "Only on service or quota failure"),
    ]
    for cell, (label, value, note) in zip(table.rows[0].cells, facts, strict=True):
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        a = p.add_run(label.upper())
        a.bold = True
        a.font.size = Pt(7.5)
        a.font.color.rgb = RGBColor.from_string(TEAL)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        a = p.add_run(value)
        a.bold = True
        a.font.size = Pt(10.3)
        a.font.color.rgb = RGBColor.from_string(INK)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        a = p.add_run(note)
        a.font.size = Pt(8.1)
        a.font.color.rgb = RGBColor.from_string("52666A")

    heading(doc, "What Folio proves")
    para(
        doc,
        "Folio is an agency-ready document-processing foundation. A vision model proposes structured invoice data; deterministic Python verifies calculations and routing policy; a reviewer resolves exceptions. The original extraction is never overwritten.",
    )
    para(
        doc,
        "The one sentence to remember: The model extracts. Deterministic code verifies. A human resolves exceptions.",
        10.4,
        TEAL,
        bold=True,
        space_after=6,
    )

    heading(doc, "Architecture and decision flow")
    flow = doc.add_table(rows=1, cols=5)
    flow.autofit = False
    labels = ["Upload", "Worker", "AI extraction", "Verify", "Decision"]
    details = [
        "PDF, PNG, JPG or WEBP\nprivate local storage",
        "Durable job\nseparate worker",
        "Groq JSON mode\nor local vision",
        "Pydantic + Decimal\nmath and policy",
        "Auto approve\nor human review",
    ]
    for i, cell in enumerate(flow.rows[0].cells):
        shade(cell, "F6F8F7" if i % 2 else PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(labels[i])
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string(INK)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(details[i])
        r.font.size = Pt(7.7)
        r.font.color.rgb = RGBColor.from_string("52666A")
    borders(flow)
    para(
        doc,
        "Every run records provider, model, prompt version, token usage when available, latency, correlation ID, checks, route reason and audit events. Idempotency prevents duplicate jobs; the worker makes one cloud request per job, never a hidden retry loop.",
        9.6,
        space_after=4,
    )

    heading(doc, "Technical essentials")
    tech = doc.add_table(rows=1, cols=2)
    tech.autofit = False
    rows = [
        ("API and UI", "FastAPI, Pydantic v2, React, TypeScript, Vite"),
        (
            "Data and jobs",
            "SQLite for local demo; SQLAlchemy; database-backed worker; Alembic migrations",
        ),
        (
            "Document handling",
            "PyMuPDF page rendering; file limits; MIME inspection; private local storage",
        ),
        (
            "Verification",
            "Python Decimal arithmetic with 0.01 tolerance; line, subtotal and total checks",
        ),
        (
            "Review and audit",
            "Separate human decision snapshot; immutable original extraction; append-only audit events",
        ),
        (
            "Providers",
            "Groq Qwen 3.8 27B primary; Gemini available; Ollama Qwen 3 VL 4B local recovery",
        ),
    ]
    for row_index, (key, value) in enumerate(rows):
        cells = tech.add_row().cells
        shade(cells[0], PALE if row_index % 2 == 0 else "F8FAFA")
        shade(cells[1], "FFFFFF")
        cell_text(cells[0], key, bold=True, color=INK, size=8.7)
        cell_text(cells[1], value, size=8.7)
    tech._tbl.remove(tech.rows[0]._tr)
    borders(tech)

    doc.add_page_break()
    heading(doc, "The recorded proof")
    para(
        doc,
        "Use the prepared two-page Alder Office Supply invoice. It is synthetic, deliberately contains a printed-total discrepancy, and makes the guardrail visible without manufacturing a failure live.",
    )

    scenario = doc.add_table(rows=1, cols=4)
    scenario.autofit = False
    headers = ["Extracted", "Printed total", "Calculated total", "Decision"]
    values = [
        "6 line items\nSubtotal 1,700",
        "USD 1,850",
        "USD 1,800",
        "Review required\nDifference USD 50",
    ]
    for cell, label, value in zip(scenario.rows[0].cells, headers, values, strict=True):
        shade(cell, INK)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label.upper())
        r.bold = True
        r.font.size = Pt(7.2)
        r.font.color.rgb = RGBColor(255, 255, 255)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(9.4)
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
    borders(scenario, INK)

    heading(doc, "60 second recording sequence")
    script = doc.add_table(rows=1, cols=3)
    script.autofit = False
    widths = [0.75, 2.12, 3.35]
    for cell, label, width in zip(
        script.rows[0].cells, ["Time", "Show", "Say"], widths, strict=True
    ):
        cell.width = Inches(width)
        shade(cell, TEAL)
        cell_text(cell, label, True, "FFFFFF", 8.5)
    steps = [
        (
            "0–8s",
            "Upload Alder PDF",
            "The model reads the invoice. It does not get to approve its own work.",
        ),
        (
            "8–18s",
            "Source and extracted fields",
            "This scanned, two-page invoice has six line items. The extraction is structured and traceable.",
        ),
        (
            "18–32s",
            "Total evidence and ledger",
            "The printed total is 1,850. Decimal arithmetic calculates 1,800. That 50-dollar difference blocks approval.",
        ),
        (
            "32–45s",
            "Reviewer decision",
            "An operator resolves the exception. The original extraction remains preserved.",
        ),
        (
            "45–54s",
            "Audit trail and trace",
            "Provider, model, latency, checks and human decision are recorded for engineering and operations.",
        ),
        (
            "54–60s",
            "Saved clean record",
            "The same foundation can be adapted to your client’s workflow and system of record.",
        ),
    ]
    for i, row in enumerate(steps):
        cells = script.add_row().cells
        for j, value in enumerate(row):
            shade(cells[j], "F6F8F7" if i % 2 else "FFFFFF")
            cell_text(cells[j], value, bold=(j == 0), color=INK, size=8.3)
    borders(script)

    heading(doc, "Before you press record")
    checks = [
        "Run .venv/bin/python scripts/preflight.py. It checks local readiness without using cloud quota.",
        "Use the saved Groq Alder record if you need a no-wait recording. It is visibly labeled GROQ LIVE and has a 3.197 second trace.",
        "Keep the local Ollama server running. If Groq quota or service availability fails, Folio may use one explicitly audited local recovery attempt.",
        "Say synthetic whenever you refer to prepared documents. Do not claim handwriting accuracy, a universal accuracy rate, zero live cost, or production certification.",
        "If you trim waiting, label it Processing wait shortened. Never present a fixture or local recovery as a Groq result.",
    ]
    for item in checks:
        bullet(doc, item)

    heading(doc, "Known boundaries")
    para(
        doc,
        "Folio is a high-quality sales demo and reusable integration base, not a certified accounting product. Groq Qwen 3.8 permits three page images per request. Handwritten or blurry documents should be treated as review-first until measured against a client-specific set. Model confidence is a heuristic, not proof. The provider sees documents when live cloud mode is selected; use approved data terms before client uploads.",
        9.2,
        space_after=0,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
